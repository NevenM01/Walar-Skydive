from openpyxl import load_workbook

wb = load_workbook("WALAR 2026-01-01.xlsx", data_only=False)

with open("formule.txt", "w", encoding="utf-8") as f:

    for sheet in wb.worksheets:
        f.write(f"=== SHEET: {sheet.title} ===\n")

        for row in sheet.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    line = f"{cell.coordinate} -> {cell.value}\n"
                    f.write(line)