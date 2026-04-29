"""Export ALL sheets from the WALAR Excel workbook to CSV for analysis.

This is intended for "Excel parity check" work: it exports raw values (data_only=True),
so formula cells are exported as their evaluated values as stored in the workbook.
"""

from __future__ import annotations

import csv
import json
import os
import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

REPO_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_XLSX = REPO_ROOT / "local" / "docs" / "walar-spec" / "WALAR 2026-01-01.xlsx"
OUT_DIR = REPO_ROOT / "local" / "docs" / "walar-spec" / "extracted"


def _safe_filename(name: str) -> str:
    # Keep it deterministic and Windows-safe.
    s = name.strip()
    s = re.sub(r"[<>:\"/\\\\|?*]+", "_", s)
    s = re.sub(r"\s+", " ", s)
    if not s:
        s = "sheet"
    return s


def _cell_to_jsonable(v: Any) -> Any:
    # openpyxl returns datetime/date/decimal/etc. json can't always serialize them.
    if v is None:
        return ""
    try:
        json.dumps(v)
        return v
    except TypeError:
        return str(v)


def main() -> None:
    path = Path(os.environ.get("WALAR_XLSX", str(DEFAULT_XLSX)))
    if not path.is_file():
        raise SystemExit(f"Workbook not found: {path}")

    OUT_DIR.mkdir(parents=True, exist_ok=True)

    wb = load_workbook(path, data_only=True, read_only=True)

    summary: dict[str, Any] = {
        "workbook": str(path),
        "sheet_count": len(wb.sheetnames),
        "sheets": [],
    }

    for name in wb.sheetnames:
        ws = wb[name]
        out_csv = OUT_DIR / f"{_safe_filename(name)}.csv"

        max_row = ws.max_row or 0
        max_col = ws.max_column or 0

        non_empty_rows = 0
        with out_csv.open("w", newline="", encoding="utf-8") as f:
            w = csv.writer(f, delimiter=";")
            for row in ws.iter_rows(values_only=True):
                vals = [_cell_to_jsonable(v) for v in row]
                # Trim trailing empty cells for readability.
                while vals and vals[-1] == "":
                    vals.pop()
                if any(v != "" for v in vals):
                    non_empty_rows += 1
                w.writerow(vals)

        summary["sheets"].append(
            {
                "name": name,
                "csv": str(out_csv.relative_to(REPO_ROOT)),
                "max_row": max_row,
                "max_col": max_col,
                "non_empty_rows": non_empty_rows,
            }
        )

    wb.close()

    summary_path = OUT_DIR / "workbook_summary.json"
    summary_path.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Wrote {len(wb.sheetnames)} CSVs to {OUT_DIR}")
    print(f"Wrote summary to {summary_path}")


if __name__ == "__main__":
    main()

