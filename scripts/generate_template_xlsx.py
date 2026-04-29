"""
Generate the WALAR results import Excel template with all required columns,
data validation dropdowns, column formatting, and an instructions sheet.
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
import os

OUT = os.path.join(os.path.dirname(__file__), "..", "web", "public", "walar-results-import.template.xlsx")

HEADERS = [
    # Competition info (used when auto-creating a new competition)
    ("Competition label",      "competition_unique_label", True,  "Unique label e.g. '2025 Europa Cup Italy'. Must match for all rows of one competition."),
    ("Competition rang",       "competition_rang_code",    False, "WPC / EPC / Wcup / CISM / SWCS / NC / WALAR A–F"),
    ("Competition location",   "competition_location",     False, "City or venue (used only when creating new competition)"),
    ("Competition start date", "competition_start_date",   False, "YYYY-MM-DD (used only when creating new competition)"),
    ("Competition end date",   "competition_end_date",     False, "YYYY-MM-DD (used only when creating new competition)"),
    (
        "Finished jump rounds",
        "competition_finished_jump_rounds",
        False,
        "Number of completed jump rounds (e.g. 8). Used to compute competition strength (OP).",
    ),

    # Athlete identity
    ("First name",       "first_name",       True,  "Athlete first / given name"),
    ("Last name",        "last_name",        True,  "Athlete last / family name"),
    ("Country code",     "country_code",     True,  "Country name or ISO code"),
    ("Gender",           "gender",           True,  "M or F"),
    ("Date of birth",    "date_of_birth",    False, "YYYY-MM-DD"),
    ("FAI licence",      "fai_licence",      False, "FAI ID number (helps match existing athletes)"),

    # GDPR & flags
    ("GDPR consent",          "gdpr_consent_given",      True,  "TRUE / FALSE — must be TRUE to appear on public leaderboard"),
    ("Publish full name",     "gdpr_publish_full_name",  False, "TRUE / FALSE"),
    ("National team member",  "member_national_team",    False, "TRUE / FALSE — affects WALAR rating (+2)"),
    ("WPC medalist",          "wpc_medalist",            False, "TRUE / FALSE — affects WALAR rating (+2)"),

    # Jump results (cm from center — 0 = dead center)
    ("Jump 1", "jump1_cm", False, "cm from center (0 = perfect)"),
    ("Jump 2", "jump2_cm", False, ""),
    ("Jump 3", "jump3_cm", False, ""),
    ("Jump 4", "jump4_cm", False, ""),
    ("Jump 5", "jump5_cm", False, ""),
    ("Jump 6", "jump6_cm", False, ""),
    ("Jump 7", "jump7_cm", False, ""),
    ("Jump 8", "jump8_cm", False, ""),
    ("SF cm",  "sf_cm",    False, "Semi-final cm"),
    ("F cm",   "f_cm",     False, "Final cm"),

    # Tie-break (optional). Fill only when ties are resolved with extra jumps. Not added to main total / AX.
    ("TB 1", "tb1_cm", False, "Tie-break round 1 (cm). Optional. Resolve final ranks in Excel; Place columns = official after TB."),
    ("TB 2", "tb2_cm", False, "Tie-break round 2 (cm). Optional."),
    ("TB 3", "tb3_cm", False, "Tie-break round 3 (cm). Optional."),
    ("TB 4", "tb4_cm", False, "Tie-break round 4 (cm). Optional."),
    ("TB 5", "tb5_cm", False, "Tie-break round 5 (cm). Optional."),
    ("TB 6", "tb6_cm", False, "Tie-break round 6 (cm). Optional."),

    # Placements
    ("Place M+F",    "place_overall", False, "Optional. Overall placement (all genders). If blank, derived from imported totals."),
    ("Place M",      "place_m",       False, "Optional. Male placement. If blank, derived from imported totals."),
    ("Place F",      "place_f",       False, "Optional. Female placement. If blank, derived from imported totals."),
    ("Place J",      "place_j",       False, "Optional. Junior placement. If blank, derived from imported totals."),
    ("Place MJ",     "place_mj",      False, "Optional. Junior male placement. If blank, derived from imported totals."),
    ("Place FJ",     "place_fj",      False, "Optional. Junior female placement. If blank, derived from imported totals."),
    ("Place Master", "place_master",  False, "Optional. Master placement. If blank, derived from imported totals."),

    # Meta
    ("Start number", "start_number", False, "Bib / start number"),
    ("Team",         "team",         False, "Team name"),
]

RANG_VALUES = "WPC,EPC,Wcup,CISM,SWCS,NC,WALAR A,WALAR B,WALAR C,WALAR D,WALAR E,WALAR F"

# Columns where empty cells are auto-filled on import from jump totals (see csvImportResults).
AUTOFILL_PLACE_KEYS = frozenset(
    {
        "place_overall",
        "place_m",
        "place_f",
        "place_j",
        "place_mj",
        "place_fj",
        "place_master",
    }
)

TB_ROUND_KEYS = frozenset({f"tb{i}_cm" for i in range(1, 7)})

HEADER_FILL = PatternFill(start_color="1B3A5C", end_color="1B3A5C", fill_type="solid")
# Distinct header color for auto-derivable placement columns (empty → filled on import).
HEADER_FILL_AUTOFILL = PatternFill(start_color="C2410C", end_color="C2410C", fill_type="solid")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
REQUIRED_FILL = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")
OPTIONAL_FILL = PatternFill(start_color="F8F9FA", end_color="F8F9FA", fill_type="solid")
# Hint row + sample row tint for the same columns (readable with dark text).
AUTOFILL_HINT_FILL = PatternFill(start_color="FFEDD5", end_color="FFEDD5", fill_type="solid")
# Tie-break columns (distinct from navy + orange place headers).
HEADER_FILL_TB = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")
TB_HINT_FILL = PatternFill(start_color="E0F2F1", end_color="E0F2F1", fill_type="solid")
THIN_BORDER = Border(
    left=Side(style="thin", color="CCCCCC"),
    right=Side(style="thin", color="CCCCCC"),
    top=Side(style="thin", color="CCCCCC"),
    bottom=Side(style="thin", color="CCCCCC"),
)

EXAMPLE_ROW = {
    "competition_unique_label": "2025 WALAR Cup Grobnik",
    "competition_rang_code": "WALAR B",
    "competition_location": "Grobnik",
    "competition_start_date": "2025-07-12",
    "competition_end_date": "2025-07-13",
    "competition_finished_jump_rounds": "8",
    "first_name": "Marko",
    "last_name": "Horvat",
    "country_code": "Croatia",
    "gender": "M",
    "date_of_birth": "1990-05-15",
    "fai_licence": "123456",
    "gdpr_consent_given": "TRUE",
    "gdpr_publish_full_name": "TRUE",
    "member_national_team": "TRUE",
    "wpc_medalist": "FALSE",
    "jump1_cm": "0",
    "jump2_cm": "1",
    "jump3_cm": "0",
    "jump4_cm": "2",
    "jump5_cm": "1",
    "jump6_cm": "0",
    "jump7_cm": "3",
    "jump8_cm": "0",
    "sf_cm": "1",
    "f_cm": "0",
    "tb1_cm": "",
    "tb2_cm": "",
    "tb3_cm": "",
    "tb4_cm": "",
    "tb5_cm": "",
    "tb6_cm": "",
    "place_overall": "",
    "place_m": "1",
    "place_f": "",
    "place_j": "",
    "place_mj": "",
    "place_fj": "",
    "place_master": "",
    "start_number": "15A",
    "team": "Croatia A",
}


def build_results_sheet(wb: Workbook):
    ws = wb.active
    ws.title = "Results"
    ws.sheet_properties.tabColor = "1B3A5C"

    for col_idx, (label, key, required, hint) in enumerate(HEADERS, 1):
        is_autofill_place = key in AUTOFILL_PLACE_KEYS
        is_tb = key in TB_ROUND_KEYS
        c = ws.cell(row=1, column=col_idx, value=label)
        c.font = HEADER_FONT
        if is_autofill_place:
            c.fill = HEADER_FILL_AUTOFILL
        elif is_tb:
            c.fill = HEADER_FILL_TB
        else:
            c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = THIN_BORDER

        # Row 2: hint / comment
        h = ws.cell(row=2, column=col_idx)
        h.value = ("REQUIRED. " if required else "") + hint
        h.font = Font(name="Calibri", size=9, italic=True, color="666666")
        if is_autofill_place:
            h.fill = AUTOFILL_HINT_FILL
        elif is_tb:
            h.fill = TB_HINT_FILL
        else:
            h.fill = REQUIRED_FILL if required else OPTIONAL_FILL
        h.alignment = Alignment(wrap_text=True, vertical="top")
        h.border = THIN_BORDER

        # Row 3: example data
        ex = ws.cell(row=3, column=col_idx, value=EXAMPLE_ROW.get(key, ""))
        ex.font = Font(name="Calibri", size=10)
        if is_autofill_place:
            ex.fill = AUTOFILL_HINT_FILL
        elif is_tb:
            ex.fill = TB_HINT_FILL
        else:
            ex.fill = PatternFill()
        ex.border = THIN_BORDER

        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = max(len(label) + 4, 14)

    # Data validations
    gender_dv = DataValidation(type="list", formula1='"M,F"', allow_blank=True)
    gender_dv.error = "Enter M or F"
    gender_dv.errorTitle = "Invalid Gender"
    ws.add_data_validation(gender_dv)
    gender_col = next(i for i, (_, k, *_) in enumerate(HEADERS, 1) if k == "gender")
    gender_dv.add(f"{get_column_letter(gender_col)}3:{get_column_letter(gender_col)}1000")

    bool_cols = ["gdpr_consent_given", "gdpr_publish_full_name", "member_national_team", "wpc_medalist"]
    for bk in bool_cols:
        dv = DataValidation(type="list", formula1='"TRUE,FALSE"', allow_blank=True)
        dv.error = "Enter TRUE or FALSE"
        ws.add_data_validation(dv)
        ci = next(i for i, (_, k, *_) in enumerate(HEADERS, 1) if k == bk)
        dv.add(f"{get_column_letter(ci)}3:{get_column_letter(ci)}1000")

    rang_dv = DataValidation(type="list", formula1=f'"{RANG_VALUES}"', allow_blank=True)
    rang_dv.error = "Pick a valid competition rang"
    ws.add_data_validation(rang_dv)
    rang_col = next(i for i, (_, k, *_) in enumerate(HEADERS, 1) if k == "competition_rang_code")
    rang_dv.add(f"{get_column_letter(rang_col)}3:{get_column_letter(rang_col)}1000")

    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}1"
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 40


def build_instructions_sheet(wb: Workbook):
    ws = wb.create_sheet("How to use")
    ws.sheet_properties.tabColor = "28A745"

    lines = [
        ("WALAR Results Import Template", True, 16),
        ("", False, 11),
        ("1. Fill the 'Results' sheet with one row per athlete per competition.", False, 11),
        ("2. Row 1 = headers (do not change). Row 2 = hints (delete before upload or leave — ignored).", False, 11),
        ("3. Row 3 has example data — replace or delete it.", False, 11),
        ("4. All athletes in the same competition must have the same 'Competition label'.", False, 11),
        ("5. If the competition already exists in the database, results are added/updated.", False, 11),
        ("6. If the competition does NOT exist, it is auto-created from the label + rang + dates + location + finished jump rounds.", False, 11),
        ("7. Athletes are matched by name + country. If not found, a new athlete is created.", False, 11),
        ("8. After import, WALAR points are automatically recalculated.", False, 11),
        ("", False, 11),
        ("REQUIRED columns: Competition label, First name, Last name, Country code, Gender, GDPR consent.", True, 11),
        ("Color key: navy = main fields; teal = tie-break TB1–TB6 (optional, not summed into main total / AX); orange = Place (optional auto-fill from totals if blank).", True, 11),
        ("Tie-break: use Excel to resolve ties; enter final official ranks in Place columns. TB columns are for records only — app does not recompute TB logic.", True, 11),
        ("IMPORTANT: Place columns are optional. If left blank, the importer auto-derives placements from totals per competition.", True, 11),
        ("Manual place values always win (use them to mirror official final ranks when they differ from pure totals).", False, 11),
        ("", False, 11),
        ("Accepted file formats: .xlsx (this template) or .csv (exported from Excel).", False, 11),
        ("Upload at: Admin → Results Import", False, 11),
    ]

    for i, (text, bold, size) in enumerate(lines, 1):
        c = ws.cell(row=i, column=1, value=text)
        c.font = Font(name="Calibri", bold=bold, size=size)

    ws.column_dimensions["A"].width = 100


def main():
    wb = Workbook()
    build_results_sheet(wb)
    build_instructions_sheet(wb)
    wb.save(OUT)
    print(f"Template saved to {OUT}")


if __name__ == "__main__":
    main()
